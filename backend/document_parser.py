"""Document parser voor Word (.docx) en PDF bestanden."""

import logging
from pathlib import Path

_log = logging.getLogger("app.document_parser")


def parse_docx(file_path: Path) -> tuple[str, list[dict]]:
    """
    Lees een Word-bestand in.

    Returns:
        (volledige tekst, lijst van {text: str, page: None})
    """
    from docx import Document  # noqa: PLC0415

    doc = Document(file_path)
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append({"text": text, "page": None})

    # Ook tabellen meenemen
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    paragraphs.append({"text": text, "page": None})

    full_text = "\n".join(p["text"] for p in paragraphs)
    _log.info("Gelezen Word-document: %d paragrafen", len(paragraphs))
    return full_text, paragraphs


def parse_pdf(file_path: Path) -> tuple[str, list[dict]]:
    """
    Lees een PDF-bestand in.

    Returns:
        (volledige tekst, lijst van {text: str, page: int})
    """
    import pdfplumber  # noqa: PLC0415

    passages = []
    full_text_parts = []

    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                full_text_parts.append(text)
                for paragraph in text.split("\n\n"):
                    paragraph = paragraph.strip()
                    if paragraph:
                        passages.append({"text": paragraph, "page": page_num})

    full_text = "\n".join(full_text_parts)
    _log.info(
        "Gelezen PDF: %d passages over %d pagina's",
        len(passages),
        len(pdf.pages) if passages else 0,
    )
    return full_text, passages
